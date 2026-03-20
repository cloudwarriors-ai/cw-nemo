"""
Report generation service for scheduled QA reports.

Generates weekly issue reports and other scheduled reports
for distribution via Zoom Team Chat.
"""

import base64
import io
import logging
from dataclasses import dataclass, field
from datetime import datetime
from typing import Optional

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

from ..cache import IssueCache


@dataclass
class WeeklyReport:
    """Data transfer object for weekly issue report."""

    generated_at: datetime
    total_issues: int
    high_priority: int
    unassigned: int
    stale: int
    by_repo: dict[str, int] = field(default_factory=dict)
    cache_age_seconds: float = 0.0
    has_valid_data: bool = True
    error: Optional[str] = None

    def to_dict(self) -> dict:
        """Convert to dictionary for JSON serialization."""
        return {
            "generated_at": self.generated_at.isoformat(),
            "total_issues": self.total_issues,
            "high_priority": self.high_priority,
            "unassigned": self.unassigned,
            "stale": self.stale,
            "by_repo": self.by_repo,
            "cache_age_seconds": self.cache_age_seconds,
            "has_valid_data": self.has_valid_data,
            "error": self.error,
        }


@dataclass
class ReportResult:
    """Result of report execution."""

    success: bool = False
    report: Optional[WeeklyReport] = None
    message_sent: bool = False
    warnings: list[str] = field(default_factory=list)
    error: Optional[str] = None


class ReportService:
    """
    Service for generating scheduled reports.

    Generates weekly issue reports from cached GitHub data
    for distribution to Zoom Team Chat channels.
    """

    def __init__(
        self,
        issue_cache: IssueCache,
        logger: logging.Logger = None
    ):
        """
        Initialize report service.

        Args:
            issue_cache: IssueCache instance for GitHub data
            logger: Logger instance
        """
        self.issue_cache = issue_cache
        self.logger = logger or logging.getLogger("qa_agent")

    def generate_weekly_report(self) -> WeeklyReport:
        """
        Generate weekly issue report from cached data.

        Returns:
            WeeklyReport with current issue statistics
        """
        try:
            # Check if cache has data
            if not self.issue_cache.has_data:
                self.logger.warning("Cache empty for weekly report")
                return WeeklyReport(
                    generated_at=datetime.now(),
                    total_issues=0,
                    high_priority=0,
                    unassigned=0,
                    stale=0,
                    has_valid_data=False,
                    error="GitHub data unavailable - cache empty"
                )

            # Get statistics from cache
            stats = self.issue_cache.get_stats()

            report = WeeklyReport(
                generated_at=datetime.now(),
                total_issues=stats.get("total", 0),
                high_priority=stats.get("high_priority", 0),
                unassigned=stats.get("unassigned", 0),
                stale=stats.get("stale", 0),
                by_repo=stats.get("by_repo", {}),
                cache_age_seconds=stats.get("cache_age_seconds", 0),
                has_valid_data=True,
            )

            self.logger.info(
                f"Generated weekly report: {report.total_issues} issues, "
                f"{report.high_priority} high priority"
            )

            return report

        except Exception as e:
            self.logger.error(f"Failed to generate weekly report: {e}")
            return WeeklyReport(
                generated_at=datetime.now(),
                total_issues=0,
                high_priority=0,
                unassigned=0,
                stale=0,
                has_valid_data=False,
                error=str(e)
            )

    def format_for_zoom(self, report: WeeklyReport) -> str:
        """
        Format weekly report for Zoom Team Chat (legacy simple format).

        Args:
            report: WeeklyReport to format

        Returns:
            Formatted markdown string for Zoom
        """
        if not report.has_valid_data:
            return (
                f"Weekly Issue Report - {report.generated_at.strftime('%B %d, %Y')}\n\n"
                f"Unable to generate report: {report.error or 'Data unavailable'}\n\n"
                "Please check GitHub connectivity and try again."
            )

        lines = [
            f"Weekly Issue Report - {report.generated_at.strftime('%B %d, %Y')}",
            "",
            "Summary:",
            f"  Total Open Issues: {report.total_issues}",
            f"  High Priority: {report.high_priority}",
            f"  Unassigned: {report.unassigned}",
            f"  Stale (14+ days): {report.stale}",
            "",
        ]

        # Add by-repo breakdown if available
        if report.by_repo:
            lines.append("By Repository:")
            for repo, count in sorted(report.by_repo.items(), key=lambda x: -x[1]):
                lines.append(f"  {repo}: {count}")
            lines.append("")

        # Add cache age warning if stale
        if report.cache_age_seconds > 3600:  # > 1 hour
            age_hours = report.cache_age_seconds / 3600
            lines.append(f"Note: Data is {age_hours:.1f} hours old")
            lines.append("")

        lines.append("For details, ask: 'show high priority issues' or 'show unassigned'")

        return "\n".join(lines)

    def format_as_table(self) -> str:
        """
        Format weekly report for Zoom chat (no fixed-width alignment).

        Returns:
            Formatted report for Zoom chat
        """
        now = datetime.now()
        date_str = now.strftime("%d%b%y @ %I %p EST").replace(" 0", " ")

        if not self.issue_cache.has_data:
            return (
                f"GITHUB Issues Tracker - {date_str}\n\n"
                "Unable to generate report: GitHub data unavailable\n"
                "Please check connectivity and try again."
            )

        # Get detailed stats per repo
        repo_stats = self._get_detailed_repo_stats()

        if not repo_stats:
            return (
                f"GITHUB Issues Tracker - {date_str}\n\n"
                "No issues found in tracked repositories."
            )

        # Calculate totals
        totals = {"open": 0, "high": 0, "unassigned": 0, "stale": 0}
        for repo in repo_stats:
            totals["open"] += repo["open"]
            totals["high"] += repo["high"]
            totals["unassigned"] += repo["unassigned"]
            totals["stale"] += repo["stale"]

        lines = [
            f"**GITHUB Issues Tracker** - {date_str}",
            "",
            f"**TOTALS:** {totals['open']} Open | {totals['high']} High Priority | {totals['unassigned']} Unassigned | {totals['stale']} Stale",
            "",
            "**By Repository:**",
        ]

        # Data rows - simple format
        for repo in repo_stats:
            line = f"- **{repo['name']}**: {repo['open']} open"
            extras = []
            if repo['high'] > 0:
                extras.append(f"{repo['high']} high")
            if repo['unassigned'] > 0:
                extras.append(f"{repo['unassigned']} unassigned")
            if repo['stale'] > 0:
                extras.append(f"{repo['stale']} stale")
            if extras:
                line += f" ({', '.join(extras)})"
            lines.append(line)

        # Add assignee breakdown
        assignee_stats = self._get_assignee_stats()
        if assignee_stats:
            lines.append("")
            lines.append("**By Assignee:**")
            for assignee, count in assignee_stats[:10]:  # Top 10
                lines.append(f"- {assignee}: {count}")
            if len(assignee_stats) > 10:
                lines.append(f"- ... and {len(assignee_stats) - 10} more")

        # Cache age note
        age_seconds = self.issue_cache.age_seconds
        if age_seconds > 3600:
            age_hours = age_seconds / 3600
            lines.append("")
            lines.append(f"_Data cached {age_hours:.1f} hours ago_")

        return "\n".join(lines)

    def _get_detailed_repo_stats(self) -> list[dict]:
        """Get detailed statistics per repository."""
        return self.issue_cache.get_detailed_repo_stats()

    def _get_assignee_stats(self) -> list[tuple[str, int]]:
        """Get issue counts by assignee."""
        return self.issue_cache.get_assignee_stats()

    def generate_excel_report(self) -> Optional[bytes]:
        """
        Generate an Excel report matching the QA format.

        Creates two sheets:
        - Report: Summary with repo stats, type breakdown, and assignee matrix
        - All Issues: Detailed list of all issues

        Returns:
            Excel file as bytes, or None if generation fails
        """
        if not OPENPYXL_AVAILABLE:
            self.logger.error("openpyxl not installed - cannot generate Excel")
            return None

        try:
            wb = Workbook()

            # Create both sheets
            self._create_report_sheet(wb)
            self._create_all_issues_sheet(wb)

            # Save to bytes
            buffer = io.BytesIO()
            wb.save(buffer)
            buffer.seek(0)

            self.logger.info("Generated Excel report successfully")
            return buffer.getvalue()

        except Exception as e:
            self.logger.error(f"Failed to generate Excel report: {e}")
            return None

    def _create_report_sheet(self, wb: "Workbook") -> None:
        """Create the Report summary sheet matching QA format."""
        ws = wb.active
        ws.title = "Report"

        # Styles
        title_font = Font(bold=True, size=14)
        header_font = Font(bold=True, size=11)
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font_white = Font(bold=True, size=11, color="FFFFFF")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        center_align = Alignment(horizontal='center')

        # Get data
        repo_stats = self._get_detailed_repo_stats_with_types()
        assignee_stats = self._get_assignee_stats()
        assignee_by_repo = self._get_assignee_by_repo_stats()

        now = datetime.now()
        date_str = now.strftime("%d%b%y @ %I %p EST").replace(" 0", " ").upper()

        # Row 1: Titles
        ws['A1'] = "GITHUB Issues Tracker"
        ws['A1'].font = title_font
        ws['M1'] = "Open Issues - Type"
        ws['M1'].font = title_font
        ws['S1'] = "Open Issues - Assignee"
        ws['S1'].font = title_font

        # Row 2: Date
        ws['A2'] = f"Date: {date_str}"
        ws['S2'] = f"Date: {date_str}"

        # Row 3: Headers for main summary (A-G)
        main_headers = [
            ("A", "Repository"),
            ("B", "App Name"),
            ("C", "Open"),
            ("D", "Closed"),
            ("E", "Total"),
            ("F", "% Complete"),
            ("G", "Change"),
        ]
        for col, header in main_headers:
            cell = ws[f"{col}3"]
            cell.value = header
            cell.font = header_font_white
            cell.fill = header_fill
            cell.border = thin_border
            cell.alignment = center_align

        # Row 3: Headers for type breakdown (M-P)
        type_headers = [
            ("M", "P1: Task"),
            ("N", "P2:Bug"),
            ("O", "P3: Features"),
            ("P", "Not Specified"),
        ]
        for col, header in type_headers:
            cell = ws[f"{col}3"]
            cell.value = header
            cell.font = header_font_white
            cell.fill = header_fill
            cell.border = thin_border
            cell.alignment = center_align

        # Row 3: Headers for assignee matrix (S onwards)
        ws['S3'] = "Name"
        ws['S3'].font = header_font_white
        ws['S3'].fill = header_fill
        ws['S3'].border = thin_border

        ws['T3'] = "Task Count"
        ws['T3'].font = header_font_white
        ws['T3'].fill = header_fill
        ws['T3'].border = thin_border

        # Add repo short names as assignee matrix columns (U onwards)
        # U is column 21
        repo_short_names = [r['name'] for r in repo_stats]
        for idx, name in enumerate(repo_short_names[:12]):  # Limit to 12 repos
            col_num = 21 + idx  # U=21, V=22, etc.
            col_letter = get_column_letter(col_num)
            cell = ws[f"{col_letter}3"]
            cell.value = name[:8]  # Truncate long names
            cell.font = header_font_white
            cell.fill = header_fill
            cell.border = thin_border
            cell.alignment = center_align

        # Data rows (starting at row 4)
        row = 4
        total_open = 0
        total_closed = 0
        total_tasks = 0
        total_bugs = 0
        total_features = 0
        total_not_specified = 0

        for repo in repo_stats:
            # Main columns (A-G)
            ws.cell(row=row, column=1, value=repo['full_name']).border = thin_border
            ws.cell(row=row, column=2, value=repo['name']).border = thin_border
            ws.cell(row=row, column=3, value=repo['open']).border = thin_border
            ws.cell(row=row, column=4, value=repo.get('closed', 0)).border = thin_border
            ws.cell(row=row, column=5, value=repo['open'] + repo.get('closed', 0)).border = thin_border

            # % Complete
            total = repo['open'] + repo.get('closed', 0)
            if total > 0:
                pct = repo.get('closed', 0) / total
                ws.cell(row=row, column=6, value=pct).border = thin_border
                ws.cell(row=row, column=6).number_format = '0%'
            else:
                ws.cell(row=row, column=6, value="-").border = thin_border

            # Change indicator (placeholder)
            ws.cell(row=row, column=7, value="-").border = thin_border

            # Type breakdown (M-P)
            ws.cell(row=row, column=13, value=repo.get('tasks', 0)).border = thin_border
            ws.cell(row=row, column=14, value=repo.get('bugs', 0)).border = thin_border
            ws.cell(row=row, column=15, value=repo.get('features', 0)).border = thin_border
            ws.cell(row=row, column=16, value=repo.get('not_specified', 0)).border = thin_border

            # Update totals
            total_open += repo['open']
            total_closed += repo.get('closed', 0)
            total_tasks += repo.get('tasks', 0)
            total_bugs += repo.get('bugs', 0)
            total_features += repo.get('features', 0)
            total_not_specified += repo.get('not_specified', 0)

            row += 1

        # Totals row
        totals_row = row
        ws.cell(row=totals_row, column=1, value="TOTAL").font = header_font
        ws.cell(row=totals_row, column=1).border = thin_border
        ws.cell(row=totals_row, column=2, value="").border = thin_border
        ws.cell(row=totals_row, column=3, value=total_open).font = header_font
        ws.cell(row=totals_row, column=3).border = thin_border
        ws.cell(row=totals_row, column=4, value=total_closed).font = header_font
        ws.cell(row=totals_row, column=4).border = thin_border
        ws.cell(row=totals_row, column=5, value=total_open + total_closed).font = header_font
        ws.cell(row=totals_row, column=5).border = thin_border

        # Total % Complete
        grand_total = total_open + total_closed
        if grand_total > 0:
            ws.cell(row=totals_row, column=6, value=total_closed / grand_total).font = header_font
            ws.cell(row=totals_row, column=6).number_format = '0%'
        else:
            ws.cell(row=totals_row, column=6, value="-").font = header_font
        ws.cell(row=totals_row, column=6).border = thin_border
        ws.cell(row=totals_row, column=7, value="").border = thin_border

        # Type totals
        ws.cell(row=totals_row, column=13, value=total_tasks).font = header_font
        ws.cell(row=totals_row, column=13).border = thin_border
        ws.cell(row=totals_row, column=14, value=total_bugs).font = header_font
        ws.cell(row=totals_row, column=14).border = thin_border
        ws.cell(row=totals_row, column=15, value=total_features).font = header_font
        ws.cell(row=totals_row, column=15).border = thin_border
        ws.cell(row=totals_row, column=16, value=total_not_specified).font = header_font
        ws.cell(row=totals_row, column=16).border = thin_border

        # Assignee matrix (S onwards, starting at row 4)
        assignee_row = 4
        for assignee, total_count in assignee_stats[:20]:  # Top 20 assignees
            ws.cell(row=assignee_row, column=19, value=assignee).border = thin_border  # S
            ws.cell(row=assignee_row, column=20, value=total_count).border = thin_border  # T

            # Per-repo counts for this assignee
            for idx, repo in enumerate(repo_stats[:12]):
                col_num = 21 + idx  # U onwards
                count = assignee_by_repo.get((assignee, repo['full_name']), 0)
                if count > 0:
                    ws.cell(row=assignee_row, column=col_num, value=count).border = thin_border
                else:
                    ws.cell(row=assignee_row, column=col_num, value="").border = thin_border

            assignee_row += 1

        # Column widths
        ws.column_dimensions['A'].width = 35
        ws.column_dimensions['B'].width = 22
        ws.column_dimensions['C'].width = 8
        ws.column_dimensions['D'].width = 8
        ws.column_dimensions['E'].width = 8
        ws.column_dimensions['F'].width = 12
        ws.column_dimensions['G'].width = 10
        ws.column_dimensions['M'].width = 10
        ws.column_dimensions['N'].width = 10
        ws.column_dimensions['O'].width = 12
        ws.column_dimensions['P'].width = 14
        ws.column_dimensions['S'].width = 18
        ws.column_dimensions['T'].width = 12

    def _create_all_issues_sheet(self, wb: "Workbook") -> None:
        """Create the All Issues detail sheet matching QA format."""
        ws = wb.create_sheet("All Issues")

        # Styles
        header_font_white = Font(bold=True, size=11, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Headers matching QA format
        headers = [
            "Repository", "System", "Issue Number", "Title", "State",
            "Created At", "Created At", "Created By", "Last Updated At",
            "Last Updated At", "Closed At", "Closed At", "Type", "Labels",
            "Assignees", "Comments Count", "URL", "Priority", "Skill"
        ]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font_white
            cell.fill = header_fill
            cell.border = thin_border

        # Get all issues from cache
        all_issues = self._get_all_issues()

        # Data rows
        row = 2
        for issue in all_issues:
            # Repository (full name)
            ws.cell(row=row, column=1, value=issue.get('repo', '')).border = thin_border

            # System (short name / app name)
            repo_name = issue.get('repo', '')
            short_name = repo_name.split('/')[-1] if '/' in repo_name else repo_name
            ws.cell(row=row, column=2, value=short_name).border = thin_border

            # Issue Number
            ws.cell(row=row, column=3, value=issue.get('number', '')).border = thin_border

            # Title
            ws.cell(row=row, column=4, value=issue.get('title', '')).border = thin_border

            # State
            ws.cell(row=row, column=5, value='open').border = thin_border

            # Created At (raw)
            created = issue.get('created_at', '')
            ws.cell(row=row, column=6, value=created).border = thin_border

            # Created At (formatted)
            formatted_created = self._format_date(created)
            ws.cell(row=row, column=7, value=formatted_created).border = thin_border

            # Created By (not available in our data)
            ws.cell(row=row, column=8, value='').border = thin_border

            # Last Updated At (raw)
            updated = issue.get('updated_at', '')
            ws.cell(row=row, column=9, value=updated).border = thin_border

            # Last Updated At (formatted)
            formatted_updated = self._format_date(updated)
            ws.cell(row=row, column=10, value=formatted_updated).border = thin_border

            # Closed At (raw) - not applicable for open issues
            ws.cell(row=row, column=11, value='').border = thin_border

            # Closed At (formatted)
            ws.cell(row=row, column=12, value='').border = thin_border

            # Type (from labels)
            issue_type = self._determine_issue_type(issue.get('labels', []))
            ws.cell(row=row, column=13, value=issue_type).border = thin_border

            # Labels (comma-separated)
            labels = issue.get('labels', [])
            labels_str = ', '.join(labels) if labels else 'None'
            ws.cell(row=row, column=14, value=labels_str).border = thin_border

            # Assignees
            assignee = issue.get('assignee', '')
            ws.cell(row=row, column=15, value=assignee or 'Unassigned').border = thin_border

            # Comments Count (not available)
            ws.cell(row=row, column=16, value='').border = thin_border

            # URL
            ws.cell(row=row, column=17, value=issue.get('url', '')).border = thin_border

            # Priority
            ws.cell(row=row, column=18, value=issue.get('priority', '')).border = thin_border

            # Skill (not available)
            ws.cell(row=row, column=19, value='').border = thin_border

            row += 1

        # Column widths
        ws.column_dimensions['A'].width = 35
        ws.column_dimensions['B'].width = 22
        ws.column_dimensions['C'].width = 12
        ws.column_dimensions['D'].width = 50
        ws.column_dimensions['E'].width = 8
        ws.column_dimensions['F'].width = 22
        ws.column_dimensions['G'].width = 12
        ws.column_dimensions['H'].width = 15
        ws.column_dimensions['I'].width = 22
        ws.column_dimensions['J'].width = 12
        ws.column_dimensions['K'].width = 22
        ws.column_dimensions['L'].width = 12
        ws.column_dimensions['M'].width = 14
        ws.column_dimensions['N'].width = 30
        ws.column_dimensions['O'].width = 18
        ws.column_dimensions['P'].width = 12
        ws.column_dimensions['Q'].width = 60
        ws.column_dimensions['R'].width = 10
        ws.column_dimensions['S'].width = 10

    def _get_detailed_repo_stats_with_types(self) -> list[dict]:
        """Get detailed statistics per repository including type breakdown."""
        return self.issue_cache.get_detailed_repo_stats_with_types(self._determine_issue_type)

    def _get_assignee_by_repo_stats(self) -> dict[tuple[str, str], int]:
        """Get issue counts by (assignee, repo) pair for the matrix."""
        return self.issue_cache.get_assignee_by_repo_stats()

    def _get_all_issues(self) -> list[dict]:
        """Get all issues from cache as dictionaries."""
        issues = self.issue_cache.get_all_issues_dict()
        # Map issue_number to number for backward compatibility
        for issue in issues:
            issue["number"] = issue.pop("issue_number")
        return issues

    def _determine_issue_type(self, labels: list) -> str:
        """Determine issue type from labels (Task, Bug, Feature, or Not Specified)."""
        if not labels:
            return "Not Specified"

        labels_lower = [l.lower() for l in labels]

        # Check for common type indicators
        for label in labels_lower:
            if 'bug' in label:
                return 'Bug'
            if 'feature' in label or 'enhancement' in label:
                return 'Feature'
            if 'task' in label:
                return 'Task'

        return "Not Specified"

    def _format_date(self, date_str: str) -> str:
        """Format ISO date string to dd-mmm-yy format."""
        if not date_str:
            return ""
        try:
            # Parse ISO format
            dt = datetime.fromisoformat(date_str.replace('Z', '+00:00'))
            return dt.strftime("%d-%b-%y")
        except (ValueError, TypeError) as e:
            self.logger.debug(f"Failed to parse date '{date_str}': {e}")
            return date_str[:10] if len(date_str) >= 10 else date_str

    def generate_excel_base64(self) -> Optional[str]:
        """
        Generate Excel report and return as base64 string.

        Returns:
            Base64-encoded Excel file, or None if generation fails
        """
        excel_bytes = self.generate_excel_report()
        if excel_bytes:
            return base64.b64encode(excel_bytes).decode('utf-8')
        return None
